#!/usr/bin/env python3
"""
Generate the UAT sign-off workbook (docs/uat/UAT-Signoff.xlsx) for the
American Club website. Content-type list is read live from the Strapi schemas
so it stays in sync; the test cases are maintained inline below.

Run:  python3 docs/uat/generate_uat_signoff.py
Deps: openpyxl
"""
import glob
import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUT = os.path.join(ROOT, "docs", "uat", "UAT-Signoff.xlsx")

NAVY = "14213D"
RED = "B3202C"
LIGHT = "EEF1F6"
GREY = "5C6470"
WHITE = "FFFFFF"
HDR_FILL = PatternFill("solid", fgColor=NAVY)
SUB_FILL = PatternFill("solid", fgColor=LIGHT)
HDR_FONT = Font(bold=True, color=WHITE, size=10)
TITLE_FONT = Font(bold=True, color=NAVY, size=15)
NOTE_FONT = Font(italic=True, color=GREY, size=9)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")
CENTER = Alignment(horizontal="center", vertical="top")
THIN = Side(style="thin", color="D5DAE2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
STATUS_LIST = '"Pass,Fail,Blocked,N/A,Not tested"'

# ── content types from the live schemas ───────────────────────────
def load_types():
    cols, sing = [], []
    for f in glob.glob(os.path.join(ROOT, "cms/src/api/*/content-types/*/schema.json")):
        s = json.load(open(f))
        info, attrs = s.get("info", {}), s.get("attributes", {})
        rec = {
            "kind": s.get("kind"),
            "name": info.get("displayName") or info.get("singularName"),
            "singular": info.get("singularName"),
            "plural": info.get("pluralName"),
            "fields": len(attrs),
        }
        (cols if rec["kind"] == "collectionType" else sing).append(rec)
    cols.sort(key=lambda r: r["name"])
    sing.sort(key=lambda r: r["name"])
    return cols, sing

# Public page a tester should open to confirm the content shows.
URL_MAP = {
    # collections
    "committee-members": "/about (General Committee section)",
    "dining-promotions": "/dining/dining-promotion",
    "restaurants": "/dining/grillhouse (any restaurant)",
    "event-spaces": "/event-spaces/the-gallbrainth-ballroom (any venue)",
    "faq-categories": "/faq",
    "faq-items": "/faq",
    "aquatics-coaches": "/fitness/aquatics (team grid) + /coaches/aquatics/<slug>",
    "fitness-facilities": "/fitness/tennis (any facility)",
    "gym-trainers": "/fitness/gym (team grid)",
    "pilates-instructors": "/fitness/pilates (team grid)",
    "tennis-coaches": "/fitness/tennis (team grid)",
    "gallery-albums": "/home-sub/gallery",
    "testimonials": "/home + detail pages (member quotes)",
    "kids-experiences": "/kids/camps (any experience)",
    "news-articles": "/home-sub/news + /home-sub/club-news/<slug>",
    "event-categories": "/whats-on (filter chips)",
    "events": "/whats-on + /whats-on/<slug>",
    # singles
    "advertise-with-us-page": "/home-sub/advertise-with-us",
    "about-page": "/about",
    "contact-us-page": "/home-sub/contact-us",
    "dining-page": "/dining",
    "dining-promotions-page": "/dining/dining-promotion",
    "event-spaces-page": "/event-spaces",
    "faq-page": "/faq",
    "fitness-page": "/fitness",
    "gallery-page": "/home-sub/gallery",
    "footer": "every page (footer)",
    "header": "every page (top nav)",
    "site-config": "Google Tag / analytics (see GA sheet) — no visible page",
    "home-page": "/home",
    "kids-page": "/kids",
    "joining-fees-page": "/membership/joining-fees",
    "niche-group-membership-page": "/membership/niche-group-membership",
    "membership-page": "/membership",
    "reciprocal-clubs-page": "/membership/reciprocal-clubs",
    "referral-page": "/membership/referal",
    "start-application-page": "/membership/start-application",
    "news-page": "/home-sub/news",
    "whats-on-page": "/whats-on",
}

wb = Workbook()

def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER

def add_sheet(title, intro, headers, rows, widths, status_cols):
    ws = wb.create_sheet(title[:31])
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    t = ws.cell(row=1, column=1, value=title)
    t.font = TITLE_FONT
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    n = ws.cell(row=2, column=1, value=intro)
    n.font = NOTE_FONT
    n.alignment = WRAP
    ws.row_dimensions[2].height = 30
    hrow = 3
    for i, h in enumerate(headers, 1):
        ws.cell(row=hrow, column=i, value=h)
    style_header(ws, hrow, len(headers))
    ws.freeze_panes = ws.cell(row=hrow + 1, column=1)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    r = hrow + 1
    for row in rows:
        for i, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=i, value=val)
            cell.alignment = WRAP if i not in status_cols else CENTER
            cell.border = BORDER
            cell.font = Font(size=9)
        if (r - hrow) % 2 == 0:
            for i in range(1, len(headers) + 1):
                if ws.cell(row=r, column=i).fill.fgColor.rgb in (None, "00000000"):
                    ws.cell(row=r, column=i).fill = SUB_FILL
        r += 1
    # status dropdowns
    dv = DataValidation(type="list", formula1=STATUS_LIST, allow_blank=True)
    ws.add_data_validation(dv)
    for sc in status_cols:
        col = get_column_letter(sc)
        dv.add(f"{col}{hrow+1}:{col}{r-1}")
    return ws

# ── Cover sheet ───────────────────────────────────────────────────
cover = wb.active
cover.title = "Sign-off"
cover.sheet_view.showGridLines = False
cover.column_dimensions["A"].width = 26
cover.column_dimensions["B"].width = 70
cover.merge_cells("A1:B1")
cover["A1"] = "The American Club — Website UAT Sign-off"
cover["A1"].font = Font(bold=True, color=NAVY, size=16)
meta = [
    ("Environment", "UAT — https://uat.amclub.org.sg  (admin: /admin)"),
    ("Build / commit", "______________________  (git short SHA on uat)"),
    ("Tester name", "______________________"),
    ("Role", "______________________"),
    ("Test start date", "______________________"),
    ("Sign-off date", "______________________"),
    ("Overall result", "PASS / PASS-WITH-NOTES / FAIL"),
    ("Approved by (name)", "______________________"),
    ("Signature", "______________________"),
]
r = 3
for k, v in meta:
    cover.cell(row=r, column=1, value=k).font = Font(bold=True, color=NAVY, size=10)
    cover.cell(row=r, column=2, value=v).font = Font(size=10)
    r += 1
r += 1
cover.cell(row=r, column=1, value="How to use").font = Font(bold=True, color=RED, size=11); r += 1
for line in [
    "• Work through each tab. Set Status to Pass / Fail / Blocked / N/A / Not tested (dropdown).",
    "• Put the reproduction + screenshot link in Notes for any Fail/Blocked.",
    "• 'Just test 1 representative entry' per collection is sufficient for field checks.",
    "• Test on UAT only. Do NOT sign off against dev.",
]:
    c = cover.cell(row=r, column=1, value=line); cover.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    c.font = Font(size=10); c.alignment = WRAP; r += 1
r += 1
cover.cell(row=r, column=1, value="Status legend").font = Font(bold=True, color=RED, size=11); r += 1
for line in [
    "Pass — works as expected.   Fail — defect (log in Notes).   Blocked — cannot test (dependency).",
    "N/A — not applicable.   Not tested — pending.",
]:
    c = cover.cell(row=r, column=1, value=line); cover.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    c.font = Font(size=10); c.alignment = WRAP; r += 1

cols, sing = load_types()

# ── 1. Content editability ────────────────────────────────────────
ce_headers = ["#", "Type", "CMS name (Content Manager)", "Public page to verify", "Fields",
              "Test steps", "Expected", "Status", "Tester", "Date", "Notes"]
ce_rows = []
i = 1
STEP = ("Content Manager → open a representative entry → edit every field → Save → Publish → "
        "open the public page and hard-refresh.")
EXP = "Every field is editable; Save + Publish succeed; the change shows on the public page."
for grp, items in (("Collection", cols), ("Single", sing)):
    for rec in items:
        url = URL_MAP.get(rec["plural"]) or URL_MAP.get(rec["singular"]) or "—"
        ce_rows.append([i, grp, rec["name"], url,
                        rec["fields"], STEP, EXP, "", "", "", ""])
        i += 1
add_sheet("1. Content Editability",
          f"One row per content type ({len(cols)} collections + {len(sing)} singles = {len(cols)+len(sing)}). "
          "For collections, testing one representative entry is enough.",
          ce_headers, ce_rows, [4, 11, 30, 34, 6, 40, 34, 11, 10, 12, 30],
          status_cols=[8])

# ── 2. Responsive ─────────────────────────────────────────────────
resp_pages = [
    ("Home", "/home"), ("About", "/about"), ("Dining (landing)", "/dining"),
    ("Dining detail", "/dining/grillhouse"), ("Dining promotions", "/dining/dining-promotion"),
    ("Fitness (landing)", "/fitness"), ("Fitness detail", "/fitness/tennis"),
    ("Kids (landing)", "/kids"), ("Kids detail", "/kids/camps"),
    ("Event Spaces (landing)", "/event-spaces"), ("Event Space detail", "/event-spaces/the-gallbrainth-ballroom"),
    ("Membership (landing)", "/membership"), ("Membership sub (joining fees)", "/membership/joining-fees"),
    ("What's On (landing)", "/whats-on"), ("Event detail", "/whats-on/<any-event>"),
    ("News", "/home-sub/news"), ("News article", "/home-sub/club-news/<any>"),
    ("Gallery", "/home-sub/gallery"), ("Contact", "/home-sub/contact-us"), ("FAQ", "/faq"),
]
resp_headers = ["#", "Page", "URL", "Desktop XL ≥1440", "Desktop 1200–1439", "Tablet 768–1199",
                "Mobile <768", "Notes"]
resp_rows = [[i + 1, p[0], p[1], "", "", "", "", ""] for i, p in enumerate(resp_pages)]
add_sheet("2. Responsive",
          "Check at the 4 project breakpoints: no horizontal overflow; images/text scale; nav collapses to "
          "hamburger on mobile and the menu works; CTAs are tappable; hero/cards reflow correctly. Status per size.",
          resp_headers, resp_rows, [4, 26, 34, 16, 16, 16, 14, 34], status_cols=[4, 5, 6, 7])

def simple_sheet(title, intro, cases, widths=None):
    headers = ["#", "Test Case", "Steps / How to test", "Expected Result", "Status", "Tester", "Date", "Notes"]
    rows = [[i + 1, c[0], c[1], c[2], "", "", "", ""] for i, c in enumerate(cases)]
    add_sheet(title, intro, headers, rows, widths or [4, 34, 46, 40, 11, 10, 12, 30], status_cols=[5])

# ── 3. Google Analytics ───────────────────────────────────────────
simple_sheet("3. Google Analytics",
    "Site Configuration → Google Tag ID accepts gtag (G-/GT-/AW-) and GTM- (auto-detected). Verify tracking.",
    [
     ("Set Google Tag ID", "Admin → Global: Site Configuration → Google Tag ID → enter the real ID (G-/GT-/AW- or GTM-) → Save + Publish.", "Field saves; value persists on reload."),
     ("Tag script loads", "Open the public site → DevTools → Network → reload.", "gtag/js (for G-/GT-/AW-) or gtm.js (for GTM-) request is made with the configured ID."),
     ("Realtime hit registered", "In GA4/GTM realtime, browse the site.", "Your session/page_view appears in realtime."),
     ("SPA page_view on navigation", "Navigate between pages without full reload.", "A page_view (gtag) / dataLayer page_view (GTM) fires on each route change."),
     ("Disable analytics", "Clear the Google Tag ID → Save + Publish → reload site.", "No analytics script loads; no tracking."),
    ])

# ── 4. SSO & User Management ───────────────────────────────────────
simple_sheet("4. SSO & User Mgmt",
    "Microsoft Entra SSO + password login + role separation. Done in /admin.",
    [
     ("Login via Microsoft SSO", "/admin → 'Continue with Microsoft' → sign in with a Club account.", "Logs in to the admin; lands on dashboard."),
     ("Login via email + password", "/admin → enter email + password → Login.", "Logs in successfully."),
     ("Forgot password", "/admin → 'Forgot your password?' → follow email.", "Reset email arrives; new password works."),
     ("Invite new user", "Settings → Users → Invite new user → assign role → send.", "Invite created (Inactive); registration link works."),
     ("Editor role restrictions", "Log in as an Editor.", "Can create/edit/publish content; CANNOT see Content-Type Builder / change Settings/users."),
     ("Super Admin access", "Log in as Super Admin.", "Full access to Settings, Users, API Tokens, etc."),
     ("Deactivate / remove user", "Settings → Users → deactivate or delete a test user.", "That user can no longer log in."),
     ("SSO whitelist (if enabled)", "Confirm only allowed domains/emails can SSO in.", "Non-whitelisted accounts are rejected; no lock-out of admins."),
    ])

# ── 5. Image / Document upload ────────────────────────────────────
simple_sheet("5. Media Upload",
    "Media Library upload + assignment + serving. Azure Blob via the folder-aware provider.",
    [
     ("Upload an image", "Media Library → Add new assets → upload a JPG/PNG.", "Uploads; thumbnail renders in the Media Library."),
     ("Upload a PDF / document", "Media Library → upload a PDF/doc.", "Uploads and is listed."),
     ("Assign image to an entry", "Open an entry → image field → pick the upload → Save + Publish.", "Image shows on the public page."),
     ("Assign PDF to a CTA / download", "Set a CTA/download href to the uploaded /uploads/...pdf → publish.", "Link opens the PDF (new tab) on the public page."),
     ("Replace an asset", "Media Library → open asset → replace file.", "New file shows everywhere it is used."),
     ("Admin thumbnails (CORS)", "Browse the Media Library grid.", "Thumbnails load (no broken images / CORS errors)."),
     ("Folder organisation", "Check seeded uploads sit under their folders (e.g. fitness/, dining/).", "Assets are organised, not all flat under 'API Uploads'."),
     ("Delete an asset", "Delete a test asset.", "Removed from library; references handled."),
     ("File naming guidance", "Upload uses lowercase-with-hyphens names.", "Names follow the convention (no spaces/caps)."),
    ])

# ── 6. Chatbot (ElevenLabs) ───────────────────────────────────────
simple_sheet("6. Chatbot (ElevenLabs)",
    "Publish-time KB sync + answer accuracy. Tone/manner is optional (fine-tuning).",
    [
     ("KB sync on publish", "Publish an allow-listed entry (e.g. an Event) with distinctive text.", "Entry is upserted to the ElevenLabs KB (doc prefixed 'am-club:')."),
     ("Answer accuracy", "Open the on-site chatbot → ask about that content.", "Bot answers correctly using the published info."),
     ("Expired content removed", "Let an event pass its date (or set expiredAt) → wait for hourly sweep.", "Bot no longer cites the expired event."),
     ("Manual re-sync", "On a published entry use 'Sync to ElevenLabs'.", "Forces an immediate KB upsert; reflected in answers."),
     ("Widget loads", "Public site → chatbot widget appears and opens.", "Agent responds to a simple greeting."),
     ("Tone & manner (optional)", "Ask a few brand questions.", "Replies match the Club's voice (optional — fine-tuning)."),
    ])

# ── 7. Content workflow & rendering ───────────────────────────────
simple_sheet("7. Workflow & Rendering",
    "Draft/publish lifecycle, clone, expiry, markdown, and CTA variants.",
    [
     ("Save draft is private", "Edit an entry → Save (not Publish).", "Change is NOT visible on the public site."),
     ("Publish goes live", "Publish the entry.", "Change appears on the public page after refresh."),
     ("Modified state", "Edit a published entry → Save (don't publish).", "Public still shows the OLD version until you publish."),
     ("Unpublish", "Entry → ⋯ → Unpublish.", "Entry disappears from the public site."),
     ("Clone entry", "Open a collection entry → Clone entry.", "Creates a draft copy with '(copy)' title and '-copy' slug."),
     ("Event/promo auto-expiry", "Confirm a past-dated event/promo.", "Hidden from listings (still reachable by direct link)."),
     ("Markdown formatting", "In a richtext field use headings/bold/lists.", "Renders correctly on the page."),
     ("Markdown link", "Add [text](https://example.com) and [text](/page).", "External opens new tab; internal navigates in-app."),
     ("Markdown PDF link", "Add [Download](/uploads/.../file.pdf).", "Opens the PDF in a NEW TAB (not blank/SPA-intercepted)."),
     ("CTA variant", "Set a CTA Variant = secondary / accent / outline / text.", "Button style changes accordingly (primary/blank = white pill)."),
     ("CTA icon & bordered", "Set a CTA icon + bordered flag.", "Trailing icon + border render as chosen."),
    ])

# ── 8. Navigation & SEO ───────────────────────────────────────────
simple_sheet("8. Navigation & SEO",
    "Header/footer/breadcrumbs and per-page SEO. (Forms are out of scope.)",
    [
     ("Header nav (desktop)", "Click every top-nav item + dropdown.", "All navigate to the correct page."),
     ("Mobile hamburger menu", "On mobile, open the menu.", "Opens; all links work; closes correctly."),
     ("Footer links", "Click footer columns, contact details, social icons.", "All resolve to correct pages/targets."),
     ("Breadcrumbs", "On a detail page, use the breadcrumb / 'Back to ...'.", "Correct labels; link to the parent section."),
     ("Logo → home", "Click the header logo.", "Returns to the home page."),
     ("Page <title>", "Check the browser tab title on several pages.", "Each page has a correct, distinct title (from CMS SEO)."),
     ("Meta description", "View source / DevTools on key pages.", "meta description present and page-appropriate."),
     ("OG / social preview", "Paste a page URL into a link-preview tool.", "Title/description/image preview render."),
     ("Favicon", "Check the browser tab icon.", "Club favicon shows."),
     ("404 / unknown route", "Visit a non-existent URL.", "Graceful not-found / redirect, not a crash."),
    ])

# ── 9. Cross-browser & performance ────────────────────────────────
simple_sheet("9. Cross-browser & Perf",
    "Render parity across browsers + basic performance. (Accessibility deferred.)",
    [
     ("Chrome (desktop)", "Load key pages in Chrome.", "Renders correctly; no console errors."),
     ("Safari (desktop)", "Load key pages in Safari.", "Renders correctly."),
     ("Edge (desktop)", "Load key pages in Edge.", "Renders correctly."),
     ("Mobile Safari (iOS)", "Load on an iPhone.", "Renders + interactions work."),
     ("Mobile Chrome (Android)", "Load on Android.", "Renders + interactions work."),
     ("Lighthouse performance", "Run Lighthouse on Home + a detail page.", "Performance score within the agreed target."),
     ("Load time / LCP", "Observe first load.", "Largest content paints quickly; no long blank screen."),
     ("No console errors", "Browse with DevTools console open.", "No uncaught errors / failed requests."),
    ])

# ── 10. Infra & security ──────────────────────────────────────────
simple_sheet("10. Infra & Security",
    "Domain/TLS, access control, environment correctness, backups.",
    [
     ("Custom domain + HTTPS", "Visit the UAT URL.", "Loads over HTTPS with a valid cert (Cloudflare → Azure managed cert)."),
     ("HTTP → HTTPS", "Visit http:// of the host.", "Redirects to https://."),
     ("Admin requires login", "Visit /admin while logged out.", "Redirected to login; no content exposed."),
     ("Role separation", "As Editor, attempt Settings/Content-Type Builder.", "Access denied / hidden."),
     ("No secrets exposed", "Check page source / network for tokens.", "No API tokens or secrets in client payloads."),
     ("Environment correctness", "Confirm UAT shows UAT content.", "No dev/localhost URLs or wrong-env content leaking."),
     ("Media origin / CORS", "Inspect image/PDF URLs on the site.", "Served from the site origin (/uploads); no CORS errors."),
     ("Backup / restore ready", "Confirm DB backup (PITR) + a manual dump exist.", "Recovery path verified before go-live."),
     ("Deploy pipeline", "Confirm a content/code change deploys to UAT.", "Push → CI → deploy works; revision healthy."),
    ])

wb.save(OUT)
print("wrote", OUT)
print("sheets:", wb.sheetnames)
